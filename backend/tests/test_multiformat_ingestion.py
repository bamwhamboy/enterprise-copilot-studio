"""End-to-end multi-format ingestion tests: upload -> READY, and
non-PDF -> the existing indexing pipeline.

Exercises the real HTTP upload endpoint (not the parser classes
directly -- that's test_parser_registry.py) for every format this
sprint supports, and proves the pre-existing chunking -> embedding ->
Qdrant flow works unmodified for a non-PDF document.
"""

import csv
import io

import docx
import fitz
import openpyxl
import pptx
import pytest
from httpx import AsyncClient

KS_BASE = "/api/v1/knowledge-sources"
DOC_BASE = "/api/v1/documents"


def _auth_headers(tokens: dict) -> dict:
    return {"Authorization": f"Bearer {tokens['access_token']}"}


async def _create_knowledge_source(client: AsyncClient, headers: dict, name: str) -> str:
    response = await client.post(KS_BASE, json={"name": name}, headers=headers)
    return response.json()["id"]


def _pdf_bytes(pages: list[str]) -> bytes:
    doc = fitz.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text)
    content = doc.tobytes()
    doc.close()
    return content


def _docx_bytes(paragraphs: list[str]) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pptx_bytes(slide_texts: list[str]) -> bytes:
    presentation = pptx.Presentation()
    for text in slide_texts:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = text
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _csv_bytes(headers: list[str], rows: list[list[str]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


def _xlsx_bytes(headers: list[str], rows: list[list[str]]) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def _upload(
    client: AsyncClient, headers: dict, ks_id: str, filename: str, content: bytes, mime: str
) -> dict:
    response = await client.post(
        f"{DOC_BASE}/upload",
        data={"knowledge_source_id": ks_id},
        files={"file": (filename, io.BytesIO(content), mime)},
        headers=headers,
    )
    assert response.status_code == 201, response.text
    return response.json()


# --- Upload -> READY, per format ---------------------------------------


@pytest.mark.asyncio
async def test_upload_pdf_reaches_ready(client: AsyncClient, register_and_login) -> None:
    headers = _auth_headers(await register_and_login(email="mf-pdf@example.com"))
    ks_id = await _create_knowledge_source(client, headers, "PDF Source")

    doc = await _upload(
        client, headers, ks_id, "agreement.pdf",
        _pdf_bytes(["The registration fee is $150 per year."]),
        "application/pdf",
    )

    assert doc["status"] == "indexed"
    assert doc["pages"] == 1


@pytest.mark.asyncio
async def test_upload_docx_reaches_ready(client: AsyncClient, register_and_login) -> None:
    headers = _auth_headers(await register_and_login(email="mf-docx@example.com"))
    ks_id = await _create_knowledge_source(client, headers, "DOCX Source")

    doc = await _upload(
        client, headers, ks_id, "agreement.docx",
        _docx_bytes(["The registration fee is $150 per year."]),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert doc["status"] == "indexed"
    assert doc["pages"] >= 1


@pytest.mark.asyncio
async def test_upload_txt_reaches_ready(client: AsyncClient, register_and_login) -> None:
    headers = _auth_headers(await register_and_login(email="mf-txt@example.com"))
    ks_id = await _create_knowledge_source(client, headers, "TXT Source")

    doc = await _upload(
        client, headers, ks_id, "notes.txt",
        b"The registration fee is $150 per year.",
        "text/plain",
    )

    assert doc["status"] == "indexed"
    assert doc["pages"] == 1


@pytest.mark.asyncio
async def test_upload_pptx_reaches_ready(client: AsyncClient, register_and_login) -> None:
    headers = _auth_headers(await register_and_login(email="mf-pptx@example.com"))
    ks_id = await _create_knowledge_source(client, headers, "PPTX Source")

    doc = await _upload(
        client, headers, ks_id, "overview.pptx",
        _pptx_bytes(["Registration Fees", "Late Fee Policy"]),
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )

    assert doc["status"] == "indexed"
    assert doc["pages"] == 2


@pytest.mark.asyncio
async def test_upload_csv_reaches_ready(client: AsyncClient, register_and_login) -> None:
    headers = _auth_headers(await register_and_login(email="mf-csv@example.com"))
    ks_id = await _create_knowledge_source(client, headers, "CSV Source")

    doc = await _upload(
        client, headers, ks_id, "fees.csv",
        _csv_bytes(["Name", "Amount"], [["Registration Fee", "150"], ["Late Fee", "25"]]),
        "text/csv",
    )

    assert doc["status"] == "indexed"
    assert doc["pages"] == 1


@pytest.mark.asyncio
async def test_upload_xlsx_reaches_ready(client: AsyncClient, register_and_login) -> None:
    headers = _auth_headers(await register_and_login(email="mf-xlsx@example.com"))
    ks_id = await _create_knowledge_source(client, headers, "XLSX Source")

    doc = await _upload(
        client, headers, ks_id, "fees.xlsx",
        _xlsx_bytes(["Name", "Amount"], [["Registration Fee", "150"], ["Late Fee", "25"]]),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    assert doc["status"] == "indexed"
    assert doc["pages"] == 1


# --- Extracted text preserves tabular column context ------------------------


@pytest.mark.asyncio
async def test_csv_extracted_text_repeats_columns_per_row(
    client: AsyncClient, register_and_login
) -> None:
    headers = _auth_headers(await register_and_login(email="mf-csv-text@example.com"))
    ks_id = await _create_knowledge_source(client, headers, "CSV Column Context Source")

    doc = await _upload(
        client, headers, ks_id, "fees.csv",
        _csv_bytes(["Name", "Amount"], [["Registration Fee", "150"]]),
        "text/csv",
    )

    get_response = await client.get(f"{DOC_BASE}/{doc['id']}", headers=headers)
    extracted_text_path = get_response.json().get("extracted_text_path")
    assert extracted_text_path

    with open(extracted_text_path) as f:
        text = f.read()

    # Every row line carries its own column names -- a chunk boundary
    # anywhere in this text still has both the value and what it means.
    assert "Name: Registration Fee" in text
    assert "Amount: 150" in text


# --- Non-PDF document reaches the existing indexing pipeline ----------------


@pytest.mark.asyncio
async def test_non_pdf_document_reaches_existing_indexing_pipeline(
    client: AsyncClient, register_and_login
) -> None:
    """A DOCX document must flow through the exact same, unmodified
    chunking -> embedding -> Qdrant pipeline a PDF does -- proving
    multi-format ingestion didn't require touching anything
    downstream of text extraction."""
    headers = _auth_headers(await register_and_login(email="mf-index@example.com"))
    ks_id = await _create_knowledge_source(client, headers, "DOCX Indexing Source")

    doc = await _upload(
        client, headers, ks_id, "policy.docx",
        _docx_bytes(["The registration fee is $150 per year. " * 20]),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    response = await client.post(f"/api/v1/index/{doc['id']}")
    assert response.status_code == 200
    body = response.json()
    assert body["index_status"] == "INDEXED"
    assert body["chunks_indexed"] > 0

    get_response = await client.get(f"{DOC_BASE}/{doc['id']}", headers=headers)
    get_body = get_response.json()
    assert get_body["chunks"] == body["chunks_indexed"]
    assert get_body["embeddings"] == body["chunks_indexed"]
