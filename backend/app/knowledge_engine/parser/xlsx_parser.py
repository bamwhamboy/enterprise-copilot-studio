"""XLSX parsing via openpyxl.

Same row-serialization reasoning as the CSV parser (see its module
docstring): each row is rendered with its column names repeated
inline so a chunk boundary landing mid-sheet never orphans a row from
its headers. Each sheet is additionally labeled with its name, since
one workbook can hold multiple, unrelated tables.

``page_count`` is the number of sheets -- a real, exact number for
XLSX, the same way slide count is for PPTX.
"""

from pathlib import Path

import openpyxl

from app.core.exceptions import DocumentProcessingError
from app.knowledge_engine.parser.base import ParsedDocument


def _cell_str(value: object) -> str:
    return "" if value is None else str(value)


class XlsxParser:
    """Extracts row-serialized text and sheet count from a .xlsx file."""

    extensions = frozenset({".xlsx"})
    mime_types = frozenset(
        {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
    )

    def extract(self, path: str | Path) -> ParsedDocument:
        xlsx_path = Path(path)

        try:
            workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises its own (non-typed) errors
            raise DocumentProcessingError(
                f"Could not open '{xlsx_path.name}' as an XLSX file: {exc}"
            ) from exc

        try:
            sheet_blocks: list[str] = []
            for sheet in workbook.worksheets:
                rows = [
                    [_cell_str(cell) for cell in row]
                    for row in sheet.iter_rows(values_only=True)
                ]
                if not rows:
                    continue

                headers = rows[0]
                lines = [
                    " | ".join(
                        f"{header}: {value}"
                        for header, value in zip(headers, row)
                        if header or value
                    )
                    for row in rows[1:]
                    if any(cell.strip() for cell in row)
                ]
                if lines:
                    sheet_blocks.append(f"Sheet: {sheet.title}\n" + "\n".join(lines))

            text = "\n\n".join(sheet_blocks).strip()
            page_count = max(len(workbook.worksheets), 1)

            return ParsedDocument(text=text, page_count=page_count)
        finally:
            workbook.close()
